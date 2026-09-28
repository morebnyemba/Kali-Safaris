# whatsappcrm_backend/customer_data/manifests.py
"""
Passenger manifest exports.

Two documents are produced for a tour date:

* Park entry manifest (ZimParks) — one row per passenger with the identity
  fields the parks authority checks at the gate.
* Passenger summary (operational) — headcount per booking group, used by the
  crew for seating and to reconcile park fees.

Both documents are built from the same data (`collect_manifest_groups`) so
their headcounts always agree.
"""

import logging
from collections import Counter
from io import BytesIO
from xml.sax.saxutils import escape

import openpyxl
from django.conf import settings
from django.db.models import Prefetch
from django.http import HttpResponse
from django.utils import timezone
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .models import Booking, Traveler

logger = logging.getLogger(__name__)

# Bookings that are actually travelling. Pending, pending-manual, cancelled and
# refunded bookings (including abandoned website checkout drafts) never appear
# on a manifest.
MANIFEST_PAYMENT_STATUSES = (Booking.PaymentStatus.PAID, Booking.PaymentStatus.DEPOSIT_PAID)

# --- Palette ---
PRIMARY = colors.HexColor('#1F3D2B')
ACCENT = colors.HexColor('#E8600A')
TINT = colors.HexColor('#F4F1EA')
ROW_ALT = colors.HexColor('#FAF8F3')
GRID = colors.HexColor('#D9D4C7')
MUTED = colors.HexColor('#6B6B6B')
TEXT = colors.HexColor('#1A1A1A')
WARN = colors.HexColor('#B45309')
OK = colors.HexColor('#15803D')

PAGE_MARGIN = 36
HEADER_BAND = 34


# --- Data collection ---

def _company_details():
    details = getattr(settings, 'COMPANY_DETAILS', None) or {}
    return {
        'name': details.get('NAME') or 'Kalai Safaris',
        'address': ", ".join(p for p in [details.get('ADDRESS_LINE_1'), details.get('ADDRESS_LINE_2')] if p),
        'phone': details.get('CONTACT_PHONE', ''),
        'email': details.get('CONTACT_EMAIL', ''),
        'website': details.get('WEBSITE', ''),
    }


def booking_group_name(booking, travelers=None):
    """Best available display name for a booking group."""
    customer = booking.customer
    if customer:
        full_name = (customer.get_full_name() or '').strip()
        if full_name:
            return full_name
        contact = getattr(customer, 'contact', None)
        if contact and (contact.name or '').strip():
            return contact.name.strip()

    payload = booking.booking_details_payload if isinstance(booking.booking_details_payload, dict) else {}
    payload_customer = payload.get('customer') if isinstance(payload.get('customer'), dict) else {}
    payload_name = str(payload_customer.get('full_name') or '').strip()
    if payload_name:
        return payload_name

    if travelers is None:
        travelers = list(booking.travelers.all())
    if travelers:
        return travelers[0].name

    return booking.booking_reference or f"Booking #{booking.pk}"


def collect_manifest_groups(booking_date):
    """
    Returns (groups, totals) for confirmed bookings on `booking_date`.

    Headcount per booking is max(booked pax, recorded travelers): a booking
    whose traveler details are only partly captured still counts every paid
    seat, and the gap is reported as "details pending" instead of silently
    shrinking the headcount.
    """
    bookings = (
        Booking.objects
        .filter(start_date=booking_date, payment_status__in=MANIFEST_PAYMENT_STATUSES)
        .select_related('customer__contact')
        .prefetch_related(Prefetch('travelers', queryset=Traveler.objects.order_by('traveler_type', 'name')))
    )

    groups = []
    for booking in bookings:
        travelers = list(booking.travelers.all())
        booked = (booking.number_of_adults or 0) + (booking.number_of_children or 0)
        recorded = len(travelers)
        headcount = max(booked, recorded)

        if recorded >= booked and recorded:
            adults = sum(1 for t in travelers if t.traveler_type == Traveler.TravelerType.ADULT)
            children = recorded - adults
        else:
            adults, children = booking.number_of_adults or 0, booking.number_of_children or 0

        groups.append({
            'booking': booking,
            'name': booking_group_name(booking, travelers),
            'reference': booking.booking_reference,
            'tour': booking.tour_name,
            'status': booking.get_payment_status_display(),
            'travelers': travelers,
            'headcount': headcount,
            'recorded': recorded,
            'missing': headcount - recorded,
            'adults': adults,
            'children': children,
            'id_docs': sum(1 for t in travelers if t.id_document),
        })

    groups.sort(key=lambda g: ((g['tour'] or '').lower(), g['name'].lower(), g['reference'] or ''))

    nationalities = Counter(
        (t.nationality or '').strip().title() or 'Not provided'
        for g in groups for t in g['travelers']
    )
    pending = sum(g['missing'] for g in groups)
    nationality_breakdown = nationalities.most_common()
    if pending:
        nationality_breakdown.append(('Details pending', pending))

    totals = {
        'bookings': len(groups),
        'passengers': sum(g['headcount'] for g in groups),
        'recorded': sum(g['recorded'] for g in groups),
        'missing': pending,
        'adults': sum(g['adults'] for g in groups),
        'children': sum(g['children'] for g in groups),
        'id_docs': sum(g['id_docs'] for g in groups),
        'tours': sorted({g['tour'] for g in groups if g['tour']}),
        'nationalities': nationality_breakdown,
    }
    return groups, totals


def summary_line(groups, totals):
    """Crew copy line, e.g. 'Chakanya(14), Nyemba(10) total: 24'."""
    parts = [f"{g['name']}({g['headcount']})" for g in groups]
    return ", ".join(parts) + f" total: {totals['passengers']}"


# --- PDF building blocks ---

def _styles():
    return {
        'title': ParagraphStyle('title', fontName='Helvetica-Bold', fontSize=18, leading=22, textColor=PRIMARY),
        'subtitle': ParagraphStyle('subtitle', fontName='Helvetica', fontSize=9.5, leading=13, textColor=MUTED),
        'label': ParagraphStyle('label', fontName='Helvetica', fontSize=7.5, leading=9, textColor=MUTED),
        'value': ParagraphStyle('value', fontName='Helvetica-Bold', fontSize=10, leading=13, textColor=TEXT),
        'kpi_value': ParagraphStyle('kpi_value', fontName='Helvetica-Bold', fontSize=18, leading=21, textColor=PRIMARY, alignment=TA_CENTER),
        'kpi_label': ParagraphStyle('kpi_label', fontName='Helvetica', fontSize=7.5, leading=9, textColor=MUTED, alignment=TA_CENTER),
        'cell': ParagraphStyle('cell', fontName='Helvetica', fontSize=8.5, leading=10.5, textColor=TEXT),
        'cell_bold': ParagraphStyle('cell_bold', fontName='Helvetica-Bold', fontSize=8.5, leading=10.5, textColor=TEXT),
        'cell_right': ParagraphStyle('cell_right', fontName='Helvetica', fontSize=8.5, leading=10.5, textColor=TEXT, alignment=TA_RIGHT),
        'pending': ParagraphStyle('pending', fontName='Helvetica-Oblique', fontSize=8.5, leading=10.5, textColor=WARN),
        'band': ParagraphStyle('band', fontName='Helvetica-Bold', fontSize=8.5, leading=11, textColor=PRIMARY),
        'section': ParagraphStyle('section', fontName='Helvetica-Bold', fontSize=11, leading=14, textColor=PRIMARY, spaceBefore=4, spaceAfter=6),
        'body': ParagraphStyle('body', fontName='Helvetica', fontSize=8.5, leading=12, textColor=TEXT),
        'copy': ParagraphStyle('copy', fontName='Courier-Bold', fontSize=9.5, leading=13, textColor=TEXT),
    }


class _NumberedCanvas(rl_canvas.Canvas):
    """Canvas that stamps 'Page X of Y' once the total page count is known."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        total = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            width, _ = self._pagesize
            self.setFont('Helvetica', 7.5)
            self.setFillColor(MUTED)
            self.drawRightString(width - PAGE_MARGIN, 16, f"Page {self._pageNumber} of {total}")
            super().showPage()
        super().save()


def _page_decorator(doc_title, booking_date):
    """Header band and footer drawn on every page."""
    company = _company_details()
    generated = timezone.localtime().strftime('%d %b %Y, %H:%M')

    def draw(canvas, doc):
        width, height = doc.pagesize
        canvas.saveState()

        canvas.setFillColor(PRIMARY)
        canvas.rect(0, height - HEADER_BAND, width, HEADER_BAND, stroke=0, fill=1)
        canvas.setFillColor(ACCENT)
        canvas.rect(0, height - HEADER_BAND - 3, width, 3, stroke=0, fill=1)

        canvas.setFillColor(colors.white)
        canvas.setFont('Helvetica-Bold', 12)
        canvas.drawString(PAGE_MARGIN, height - 22, company['name'].upper())
        canvas.setFont('Helvetica', 9)
        canvas.drawRightString(
            width - PAGE_MARGIN, height - 22,
            f"{doc_title}  |  {booking_date.strftime('%A, %d %B %Y')}",
        )

        canvas.setStrokeColor(GRID)
        canvas.setLineWidth(0.5)
        canvas.line(PAGE_MARGIN, 40, width - PAGE_MARGIN, 40)
        canvas.setFillColor(MUTED)
        canvas.setFont('Helvetica', 7.5)
        contact_bits = [b for b in [company['address'], company['phone'], company['email'], company['website']] if b]
        canvas.drawString(PAGE_MARGIN, 29, " · ".join(contact_bits) or company['name'])
        canvas.drawString(PAGE_MARGIN, 16, f"Generated {generated}  ·  Confidential — contains personal data")
        canvas.restoreState()

    return draw


def _build_pdf(elements, pagesize, doc_title, booking_date, filename):
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=pagesize,
        leftMargin=PAGE_MARGIN,
        rightMargin=PAGE_MARGIN,
        topMargin=HEADER_BAND + 22,
        bottomMargin=54,
        title=f"{doc_title} — {booking_date.isoformat()}",
        author=_company_details()['name'],
    )
    decorate = _page_decorator(doc_title, booking_date)
    doc.build(elements, onFirstPage=decorate, onLaterPages=decorate, canvasmaker=_NumberedCanvas)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def _meta_table(rows, width, styles):
    """Key/value grid: rows is a list of [(label, value), ...] lines."""
    cols = max(len(r) for r in rows)
    col_width = width / cols
    data = []
    style = [
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]
    for row_idx, row in enumerate(rows):
        cells = [
            [Paragraph(escape(label).upper(), styles['label']), Paragraph(value, styles['value'])]
            for label, value in row
        ]
        if len(cells) < cols:
            # Let the last item of a short row use the remaining width.
            style.append(('SPAN', (len(cells) - 1, row_idx), (-1, row_idx)))
            cells += [''] * (cols - len(cells))
        data.append(cells)
    table = Table(data, colWidths=[col_width] * cols, hAlign='LEFT')
    table.setStyle(TableStyle(style))
    return table


def _kpi_tiles(items, width, styles):
    gap = 8
    tile_width = (width - gap * (len(items) - 1)) / len(items)
    cells, col_widths = [], []
    for idx, (value, label) in enumerate(items):
        cells.append([Paragraph(escape(str(value)), styles['kpi_value']), Paragraph(escape(label).upper(), styles['kpi_label'])])
        col_widths.append(tile_width)
        if idx < len(items) - 1:
            cells.append('')
            col_widths.append(gap)
    table = Table([cells], colWidths=col_widths, hAlign='LEFT')
    style = [
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
    ]
    for col in range(0, len(col_widths), 2):
        style += [
            ('BACKGROUND', (col, 0), (col, 0), TINT),
            ('LINEABOVE', (col, 0), (col, 0), 2, ACCENT),
        ]
    table.setStyle(TableStyle(style))
    return table


def _empty_state(message, styles):
    return Table(
        [[Paragraph(escape(message), styles['pending'])]],
        style=TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), TINT),
            ('BOX', (0, 0), (-1, -1), 0.5, GRID),
            ('TOPPADDING', (0, 0), (-1, -1), 14),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 14),
            ('LEFTPADDING', (0, 0), (-1, -1), 12),
        ]),
    )


def _base_table_style(header_rows=1):
    return [
        ('BACKGROUND', (0, 0), (-1, header_rows - 1), PRIMARY),
        ('TEXTCOLOR', (0, 0), (-1, header_rows - 1), colors.white),
        ('FONTNAME', (0, 0), (-1, header_rows - 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8.5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('LINEBELOW', (0, 0), (-1, -1), 0.4, GRID),
        ('BOX', (0, 0), (-1, -1), 0.6, GRID),
    ]


# --- Park entry manifest (ZimParks) ---

def export_booking_manifest_pdf(booking_date):
    """
    Park entry manifest for confirmed bookings on `booking_date`, landscape,
    one numbered row per passenger grouped by booking. Passengers whose
    details have not been captured yet get a numbered placeholder row so the
    final row number always equals the headcount.
    """
    styles = _styles()
    pagesize = landscape(letter)
    width = pagesize[0] - 2 * PAGE_MARGIN
    groups, totals = collect_manifest_groups(booking_date)
    company = _company_details()

    elements = [
        Paragraph("Park Entry Passenger Manifest", styles['title']),
        Paragraph(
            f"{escape(company['name'])} · Confirmed passengers for {booking_date.strftime('%A, %d %B %Y')}",
            styles['subtitle'],
        ),
        Spacer(1, 12),
    ]

    nationality_text = " · ".join(f"{escape(n)} {c}" for n, c in totals['nationalities']) or '—'
    elements.append(_meta_table([
        [
            ('Tour date', booking_date.strftime('%d %B %Y')),
            ('Operator', escape(company['name'])),
            ('Tour(s)', escape(", ".join(totals['tours'])) or '—'),
        ],
        [('Nationalities', nationality_text)],
    ], width, styles))
    elements.append(Spacer(1, 6))
    elements.append(_kpi_tiles([
        (totals['passengers'], 'Total passengers'),
        (totals['adults'], 'Adults'),
        (totals['children'], 'Children'),
        (totals['bookings'], 'Booking groups'),
        (f"{totals['id_docs']}/{totals['passengers']}", 'ID copies on file'),
    ], width, styles))
    elements.append(Spacer(1, 14))

    if not groups:
        elements.append(_empty_state(
            f"No confirmed bookings for {booking_date.strftime('%d %B %Y')}. "
            "Only bookings marked Paid or Deposit Paid appear on the manifest.",
            styles,
        ))
        return _build_pdf(elements, pagesize, 'Park Entry Manifest', booking_date,
                          f"booking_manifest_{booking_date.isoformat()}.pdf")

    col_widths = [26, 206, 126, 118, 60, 40, 70, 74]
    header = ['#', 'Full name', 'ID / Passport no.', 'Nationality', 'Gender', 'Age', 'Category', 'ID copy']
    data = [header]
    style = _base_table_style()
    style += [
        ('ALIGN', (0, 1), (0, -1), 'RIGHT'),
        ('ALIGN', (5, 1), (5, -1), 'RIGHT'),
        ('TEXTCOLOR', (0, 1), (0, -1), MUTED),
    ]

    seq = 0
    for group in groups:
        row = len(data)
        band = (
            f"{escape(group['name'])}  ·  {escape(group['reference'] or '—')}  ·  "
            f"{escape(group['tour'] or '')}  ·  {group['headcount']} pax  ·  {escape(group['status'])}"
        )
        data.append([Paragraph(band, styles['band'])] + [''] * (len(header) - 1))
        style += [
            ('SPAN', (0, row), (-1, row)),
            ('BACKGROUND', (0, row), (-1, row), TINT),
            ('LINEABOVE', (0, row), (-1, row), 0.8, PRIMARY),
        ]

        for idx, traveler in enumerate(group['travelers']):
            seq += 1
            row = len(data)
            has_doc = bool(traveler.id_document)
            data.append([
                str(seq),
                Paragraph(escape(traveler.name or ''), styles['cell_bold']),
                traveler.id_number or 'Not provided',
                traveler.nationality or 'Not provided',
                (traveler.gender or '—').title(),
                str(traveler.age) if traveler.age is not None else '—',
                traveler.get_traveler_type_display(),
                'On file' if has_doc else 'Missing',
            ])
            style.append(('TEXTCOLOR', (7, row), (7, row), OK if has_doc else WARN))
            if idx % 2:
                style.append(('BACKGROUND', (0, row), (-1, row), ROW_ALT))
            if not traveler.id_number:
                style.append(('TEXTCOLOR', (2, row), (2, row), WARN))

        for pending_idx in range(group['missing']):
            seq += 1
            row = len(data)
            data.append([
                str(seq),
                Paragraph(
                    f"Details pending — passenger {group['recorded'] + pending_idx + 1} of {group['headcount']}",
                    styles['pending'],
                ),
                '', '', '', '', '', '',
            ])
            style += [
                ('SPAN', (1, row), (-1, row)),
                ('BACKGROUND', (0, row), (-1, row), colors.HexColor('#FFF7ED')),
            ]

    table = Table(data, colWidths=col_widths, repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle(style))
    elements.append(table)

    elements.append(Spacer(1, 10))
    notes = [
        f"<b>{totals['passengers']}</b> passenger(s) across <b>{totals['bookings']}</b> booking group(s).",
        "ID copy: <font color='#15803D'>On file</font> = ID/passport image received; "
        "<font color='#B45309'>Missing</font> = still to be collected.",
    ]
    if totals['missing']:
        notes.append(
            f"<font color='#B45309'><b>{totals['missing']} passenger(s) have no details captured yet</b></font> — "
            "collect them before departure."
        )
    elements.append(Paragraph("<br/>".join(notes), styles['body']))
    elements.append(Spacer(1, 18))
    elements.append(KeepTogether([
        Paragraph("Declaration", styles['section']),
        Paragraph(
            "I confirm that the passengers listed above are travelling under this operator on the date shown.",
            styles['body'],
        ),
        Spacer(1, 16),
        _signature_block(width, styles),
    ]))

    return _build_pdf(elements, pagesize, 'Park Entry Manifest', booking_date,
                      f"booking_manifest_{booking_date.isoformat()}.pdf")


def _signature_block(width, styles):
    roles = ['Prepared by (operator)', 'Guide / Skipper', 'Parks officer']
    gap = 24
    col = (width - gap * (len(roles) - 1)) / len(roles)
    lines, labels, col_widths = [], [], []
    for idx, role in enumerate(roles):
        lines.append('')
        labels.append(Paragraph(f"{escape(role)}<br/><font color='#6B6B6B'>Name, signature &amp; date</font>", styles['label']))
        col_widths.append(col)
        if idx < len(roles) - 1:
            lines.append('')
            labels.append('')
            col_widths.append(gap)
    table = Table([lines, labels], colWidths=col_widths, rowHeights=[26, None], hAlign='LEFT')
    style = [('LEFTPADDING', (0, 0), (-1, -1), 0), ('TOPPADDING', (0, 1), (-1, 1), 4)]
    for c in range(0, len(col_widths), 2):
        style.append(('LINEBELOW', (c, 0), (c, 0), 0.8, TEXT))
    table.setStyle(TableStyle(style))
    return table


# --- Operational passenger summary ---

def export_passenger_manifest_summary_pdf(booking_date):
    """
    Operational headcount per booking group for crew seating and park fees.
    Format of the copy line: Customer(count), Customer(count) total: X
    """
    styles = _styles()
    pagesize = letter
    width = pagesize[0] - 2 * PAGE_MARGIN
    groups, totals = collect_manifest_groups(booking_date)

    elements = [
        Paragraph("Passenger Summary", styles['title']),
        Paragraph(
            f"Operational headcount · {booking_date.strftime('%A, %d %B %Y')}",
            styles['subtitle'],
        ),
        Spacer(1, 14),
        _kpi_tiles([
            (totals['passengers'], 'Passengers'),
            (totals['bookings'], 'Groups'),
            (f"{totals['adults']} / {totals['children']}", 'Adults / children'),
            (f"{totals['id_docs']}/{totals['passengers']}", 'ID copies'),
        ], width, styles),
        Spacer(1, 16),
    ]

    if not groups:
        elements.append(_empty_state(
            f"No confirmed bookings for {booking_date.strftime('%d %B %Y')}. "
            "Only bookings marked Paid or Deposit Paid are counted.",
            styles,
        ))
        return _build_pdf(elements, pagesize, 'Passenger Summary', booking_date,
                          f"passenger_summary_{booking_date.isoformat()}.pdf")

    elements.append(Paragraph("Crew copy line", styles['section']))
    copy_box = Table(
        [[Paragraph(escape(summary_line(groups, totals)), styles['copy'])]],
        colWidths=[width],
        style=TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), TINT),
            ('LINEBEFORE', (0, 0), (0, -1), 3, ACCENT),
            ('TOPPADDING', (0, 0), (-1, -1), 8),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ]),
    )
    elements += [copy_box, Spacer(1, 16), Paragraph("Breakdown by booking group", styles['section'])]

    header = ['Group', 'Reference', 'Tour', 'Adults', 'Kids', 'Pax', 'Details', 'IDs']
    col_widths = [118, 104, 118, 40, 34, 34, 46, 46]
    data = [header]
    style = _base_table_style()
    style += [('ALIGN', (3, 0), (-1, -1), 'CENTER')]

    for idx, group in enumerate(groups, start=1):
        data.append([
            Paragraph(escape(group['name']), styles['cell_bold']),
            Paragraph(escape(group['reference'] or '—'), styles['cell']),
            Paragraph(escape(group['tour'] or '—'), styles['cell']),
            str(group['adults']),
            str(group['children']),
            str(group['headcount']),
            f"{group['recorded']}/{group['headcount']}",
            f"{group['id_docs']}/{group['headcount']}",
        ])
        if idx % 2 == 0:
            style.append(('BACKGROUND', (0, idx), (-1, idx), ROW_ALT))
        if group['missing']:
            style.append(('TEXTCOLOR', (6, idx), (6, idx), WARN))
        if group['id_docs'] < group['headcount']:
            style.append(('TEXTCOLOR', (7, idx), (7, idx), WARN))

    data.append([
        Paragraph('<b>Total</b>', styles['cell']), '', '',
        str(totals['adults']), str(totals['children']), str(totals['passengers']),
        f"{totals['recorded']}/{totals['passengers']}",
        f"{totals['id_docs']}/{totals['passengers']}",
    ])
    style += [
        ('BACKGROUND', (0, -1), (-1, -1), TINT),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('LINEABOVE', (0, -1), (-1, -1), 1, PRIMARY),
    ]
    table = Table(data, colWidths=col_widths, repeatRows=1, hAlign='LEFT')
    table.setStyle(TableStyle(style))
    elements.append(table)

    elements.append(Spacer(1, 14))
    elements.append(Paragraph(
        "<b>How to read this report</b><br/>"
        "• <b>Pax</b> — seats to allocate for the group (booked count, or recorded travelers if higher).<br/>"
        "• <b>Total</b> — use for park fees and vehicle/boat capacity.<br/>"
        "• <b>Details</b> — travelers whose name/ID details have been captured.<br/>"
        "• <b>IDs</b> — travelers with an ID/passport image on file (WhatsApp or website).",
        styles['body'],
    ))

    return _build_pdf(elements, pagesize, 'Passenger Summary', booking_date,
                      f"passenger_summary_{booking_date.isoformat()}.pdf")


# --- Excel ---

_XL_TITLE = Font(bold=True, size=15, color='1F3D2B')
_XL_SUBTLE = Font(italic=True, size=9, color='6B6B6B')
_XL_HEADER_FONT = Font(bold=True, color='FFFFFF')
_XL_HEADER_FILL = PatternFill('solid', fgColor='1F3D2B')
_XL_TINT_FILL = PatternFill('solid', fgColor='F4F1EA')
_XL_WARN_FONT = Font(italic=True, color='B45309')
_XL_BORDER = Border(bottom=Side(style='thin', color='D9D4C7'))


def _xl_response(workbook, filename):
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    workbook.save(response)
    return response


def _xl_heading(sheet, title, booking_date, span):
    company = _company_details()['name']
    sheet.cell(row=1, column=1, value=company).font = _XL_TITLE
    sheet.cell(row=2, column=1, value=f"{title} — {booking_date.strftime('%A, %d %B %Y')}").font = Font(bold=True, size=12)
    sheet.cell(row=3, column=1, value=f"Generated {timezone.localtime().strftime('%d %b %Y %H:%M')}").font = _XL_SUBTLE
    for row in (1, 2, 3):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)


def _xl_header_row(sheet, row, headers, widths):
    for col, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = sheet.cell(row=row, column=col, value=header)
        cell.font = _XL_HEADER_FONT
        cell.fill = _XL_HEADER_FILL
        cell.alignment = Alignment(vertical='center')
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.freeze_panes = sheet.cell(row=row + 1, column=1)


def export_passenger_manifest_summary_excel(booking_date):
    """Excel version of the operational passenger summary."""
    groups, totals = collect_manifest_groups(booking_date)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Passenger Summary'
    headers = ['Group', 'Booking Reference', 'Tour', 'Status', 'Adults', 'Children', 'Passengers', 'Details Captured', 'ID Docs']
    _xl_heading(sheet, 'Passenger Summary', booking_date, len(headers))

    if not groups:
        sheet.cell(row=5, column=1, value=f"No confirmed bookings found for {booking_date.strftime('%B %d, %Y')}").font = _XL_WARN_FONT
        return _xl_response(workbook, f"passenger_summary_{booking_date.isoformat()}.xlsx")

    sheet.cell(row=5, column=1, value='Crew copy line').font = Font(bold=True)
    copy_cell = sheet.cell(row=6, column=1, value=summary_line(groups, totals))
    copy_cell.fill = _XL_TINT_FILL
    sheet.merge_cells(start_row=6, start_column=1, end_row=6, end_column=len(headers))

    header_row = 8
    _xl_header_row(sheet, header_row, headers, [28, 22, 30, 16, 9, 10, 12, 16, 10])
    row = header_row + 1
    for group in groups:
        values = [
            group['name'], group['reference'], group['tour'], group['status'],
            group['adults'], group['children'], group['headcount'],
            f"{group['recorded']}/{group['headcount']}", f"{group['id_docs']}/{group['headcount']}",
        ]
        for col, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=col, value=value)
            cell.border = _XL_BORDER
        row += 1

    total_values = {
        1: 'TOTAL', 5: totals['adults'], 6: totals['children'], 7: totals['passengers'],
        8: f"{totals['recorded']}/{totals['passengers']}", 9: f"{totals['id_docs']}/{totals['passengers']}",
    }
    for col in range(1, len(headers) + 1):
        cell = sheet.cell(row=row, column=col, value=total_values.get(col))
        cell.font = Font(bold=True)
        cell.fill = _XL_TINT_FILL

    return _xl_response(workbook, f"passenger_summary_{booking_date.isoformat()}.xlsx")


def export_booking_manifest_excel(booking_date):
    """Flat, filterable Excel version of the park entry manifest."""
    groups, totals = collect_manifest_groups(booking_date)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Park Manifest'
    headers = ['#', 'Group', 'Booking Reference', 'Tour', 'Full Name', 'ID / Passport', 'Nationality',
               'Gender', 'Age', 'Category', 'ID Copy']
    _xl_heading(sheet, 'Park Entry Passenger Manifest', booking_date, len(headers))
    sheet.cell(
        row=4, column=1,
        value=f"{totals['passengers']} passenger(s) · {totals['bookings']} booking group(s) · "
              f"{totals['id_docs']} ID copies on file",
    ).font = _XL_SUBTLE

    if not groups:
        sheet.cell(row=6, column=1, value=f"No confirmed bookings found for {booking_date.strftime('%B %d, %Y')}").font = _XL_WARN_FONT
        return _xl_response(workbook, f"booking_manifest_{booking_date.isoformat()}.xlsx")

    header_row = 6
    _xl_header_row(sheet, header_row, headers, [5, 26, 22, 28, 30, 20, 18, 10, 6, 10, 10])
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row}"

    row, seq = header_row + 1, 0
    for group in groups:
        prefix = [group['name'], group['reference'], group['tour']]
        for traveler in group['travelers']:
            seq += 1
            values = [seq, *prefix, traveler.name, traveler.id_number or 'Not provided',
                      traveler.nationality or 'Not provided', (traveler.gender or '').title(), traveler.age,
                      traveler.get_traveler_type_display(), 'On file' if traveler.id_document else 'Missing']
            for col, value in enumerate(values, start=1):
                sheet.cell(row=row, column=col, value=value).border = _XL_BORDER
            row += 1
        for pending_idx in range(group['missing']):
            seq += 1
            values = [seq, *prefix,
                      f"Details pending — passenger {group['recorded'] + pending_idx + 1} of {group['headcount']}"]
            for col, value in enumerate(values, start=1):
                cell = sheet.cell(row=row, column=col, value=value)
                cell.border = _XL_BORDER
                if col == 5:
                    cell.font = _XL_WARN_FONT
            row += 1

    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{row - 1}"
    return _xl_response(workbook, f"booking_manifest_{booking_date.isoformat()}.xlsx")
